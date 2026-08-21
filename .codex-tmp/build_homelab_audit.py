from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\Pc-Leandro\Desktop\Gestor\docs\Auditoria_Homelab_2026-07-29.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5F6B7A"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E8F3EC"
PALE_AMBER = "FFF4D6"
PALE_RED = "FCE8E6"
GREEN = "1F6B45"
AMBER = "8A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


PT_DIACRITICS = {
    "acessivel": "acessível",
    "administracao": "administração",
    "alem": "além",
    "alteracao": "alteração",
    "alteracoes": "alterações",
    "aplicacao": "aplicação",
    "aprovacao": "aprovação",
    "atras": "atrás",
    "atualizacao": "atualização",
    "atualizacoes": "atualizações",
    "autenticacao": "autenticação",
    "avaliacao": "avaliação",
    "cabecalho": "cabeçalho",
    "cabecalhos": "cabeçalhos",
    "calendario": "calendário",
    "categorias": "categorias",
    "certificados": "certificados",
    "classificacao": "classificação",
    "codigo": "código",
    "condicao": "condição",
    "configuracao": "configuração",
    "configuracoes": "configurações",
    "confirmacao": "confirmação",
    "conclusao": "conclusão",
    "confiaveis": "confiáveis",
    "conteudo": "conteúdo",
    "contem": "contém",
    "correcao": "correção",
    "correcoes": "correções",
    "corrupcao": "corrupção",
    "critica": "crítica",
    "critico": "crítico",
    "criticos": "críticos",
    "criptografico": "criptográfico",
    "decisao": "decisão",
    "desabilitado": "desabilitado",
    "diretorios": "diretórios",
    "disponiveis": "disponíveis",
    "dominio": "domínio",
    "dominios": "domínios",
    "endereco": "endereço",
    "enderecos": "endereços",
    "esta": "está",
    "estao": "estão",
    "estavel": "estável",
    "evidencia": "evidência",
    "evidencias": "evidências",
    "execucao": "execução",
    "exposicao": "exposição",
    "explicita": "explícita",
    "fisico": "físico",
    "forca": "força",
    "gestao": "gestão",
    "historico": "histórico",
    "horario": "horário",
    "imagens": "imagens",
    "imediata": "imediata",
    "implementacao": "implementação",
    "informacao": "informação",
    "informacoes": "informações",
    "integra": "íntegra",
    "leitura": "leitura",
    "logico": "lógico",
    "manutencao": "manutenção",
    "media": "média",
    "memoria": "memória",
    "metodo": "método",
    "migracao": "migração",
    "minima": "mínima",
    "nao": "não",
    "necessaria": "necessária",
    "necessario": "necessário",
    "nenhum": "nenhum",
    "observacao": "observação",
    "operacao": "operação",
    "operacoes": "operações",
    "opcoes": "opções",
    "periodicas": "periódicas",
    "permissoes": "permissões",
    "persistencia": "persistência",
    "politica": "política",
    "possivel": "possível",
    "pressao": "pressão",
    "producao": "produção",
    "protecao": "proteção",
    "proximos": "próximos",
    "publica": "pública",
    "publicacao": "publicação",
    "publicas": "públicas",
    "publico": "público",
    "publicos": "públicos",
    "rapidas": "rápidas",
    "recuperacao": "recuperação",
    "recomendacao": "recomendação",
    "reducao": "redução",
    "reinicializacao": "reinicialização",
    "reinicializacoes": "reinicializações",
    "relatorio": "relatório",
    "remocao": "remoção",
    "renovacao": "renovação",
    "renovacoes": "renovações",
    "resiliencia": "resiliência",
    "resolucao": "resolução",
    "responsavel": "responsável",
    "restauracao": "restauração",
    "retencao": "retenção",
    "revisao": "revisão",
    "rotacao": "rotação",
    "seguranca": "segurança",
    "sensiveis": "sensíveis",
    "sessao": "sessão",
    "sessoes": "sessões",
    "sincronizacao": "sincronização",
    "superficie": "superfície",
    "tecnica": "técnica",
    "tecnicas": "técnicas",
    "tecnologia": "tecnologia",
    "temporario": "temporário",
    "trafego": "tráfego",
    "ultimo": "último",
    "usuarios": "usuários",
    "validacao": "validação",
    "variaveis": "variáveis",
    "verificavel": "verificável",
    "versao": "versão",
    "versoes": "versões",
    "acessiveis": "acessíveis",
    "apos": "após",
    "ausencia": "ausência",
    "automacao": "automação",
    "comparacao": "comparação",
    "copia": "cópia",
    "criterio": "critério",
    "criterios": "critérios",
    "dependencia": "dependência",
    "dependencias": "dependências",
    "dimensao": "dimensão",
    "elevacao": "elevação",
    "existencia": "existência",
    "expiracao": "expiração",
    "explicito": "explícito",
    "explicitos": "explícitos",
    "ha": "há",
    "instancia": "instância",
    "instancias": "instâncias",
    "limitacoes": "limitações",
    "logicas": "lógicas",
    "medicoes": "medições",
    "medio": "médio",
    "oportunistica": "oportunística",
    "padrao": "padrão",
    "padroes": "padrões",
    "privilegio": "privilégio",
    "privilegios": "privilégios",
    "relacao": "relação",
    "segmentacao": "segmentação",
    "separacao": "separação",
    "servico": "serviço",
    "servicos": "serviços",
    "superficies": "superfícies",
    "tambem": "também",
    "tecnico": "técnico",
    "tecnicos": "técnicos",
    "transitorio": "transitório",
    "tunel": "túnel",
    "unicos": "únicos",
    "ultimas": "últimas",
    "usuario": "usuário",
}


def _match_case(source, target):
    if source.isupper():
        return target.upper()
    if source[:1].isupper():
        return target[:1].upper() + target[1:]
    return target


def apply_portuguese_diacritics(doc):
    roots = [doc.element.body]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])
    pattern = re.compile(r"\b(" + "|".join(sorted(map(re.escape, PT_DIACRITICS), key=len, reverse=True)) + r")\b", re.IGNORECASE)
    for root in roots:
        for node in root.iter(qn("w:t")):
            if not node.text:
                continue
            text = re.sub(r"\bnão e\b", "não é", node.text, flags=re.IGNORECASE)
            text = re.sub(r"\bprioridade absoluta e estabelecer\b", "prioridade absoluta é estabelecer", text, flags=re.IGNORECASE)
            node.text = pattern.sub(lambda m: _match_case(m.group(0), PT_DIACRITICS[m.group(0).lower()]), text)


def set_cell_text(cell, text, bold=False, color=INK, size=9.3, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    set_font(r, "Arial", size, color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name="Arial", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == TABLE_WIDTH_DXA, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)


def set_table_borders(table, color="D7DCE3", size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, sep, fallback, end])
    set_font(run, "Arial", 8.5, MUTED)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_paragraph(doc, text="", bold_prefix=None, italic=False, color=INK, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, color=color, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, color=color, italic=italic)
    if keep:
        set_keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Audit Bullet")
    apply_paragraph_numbering(p, 90, level)
    r = p.add_run(text)
    set_font(r, color=INK)
    return p


def add_numbered(doc, text, level=0):
    p = doc.add_paragraph(style="Audit Number")
    apply_paragraph_numbering(p, 91, level)
    r = p.add_run(text)
    set_font(r, color=INK)
    return p


def apply_paragraph_numbering(paragraph, num_id, level=0):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def add_callout(doc, label, text, kind="info"):
    palette = {
        "info": (PALE_BLUE, DARK_BLUE),
        "positive": (PALE_GREEN, GREEN),
        "warning": (PALE_AMBER, AMBER),
        "risk": (PALE_RED, RED),
    }
    fill, accent = palette[kind]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    r1 = p.add_run(label + " ")
    set_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=10.5, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        set_cell_shading(header.cells[idx], LIGHT_GRAY)
        set_cell_text(header.cells[idx], value, bold=True, color=DARK_BLUE, size=9.1)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        if ridx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "FAFBFC")
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 and widths[idx] <= 1500 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[idx], value, color=INK, size=font_size, align=align)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    heading_tokens = {
        1: (16, BLUE, 12, 6),
        2: (13, BLUE, 10, 5),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    for name, num_id in (("Audit Bullet", 90), ("Audit Number", 91)):
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.base_style = styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        p_pr = style.element.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num])


def add_custom_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstracts = []
    nums = []
    for abstract_id, num_id, fmt, text_value, font in (
        (90, 90, "bullet", "•", "Arial"),
        (91, 91, "decimal", "%1.", "Arial"),
    ):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text_value)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), font)
        fonts.set(qn("w:hAnsi"), font)
        r_pr.append(fonts)
        lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
        abstract.append(lvl)
        abstracts.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        nums.append(num)
    for abstract in abstracts:
        numbering.append(abstract)
    for num in nums:
        numbering.append(num)


def setup_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def setup_header_footer(section, first=False):
    setup_page(section)
    section.different_first_page_header_footer = first
    if not first:
        header = section.header
        p = header.paragraphs[0]
        p.text = ""
        p.paragraph_format.space_after = Pt(0)
        left = p.add_run("AUDITORIA DE SEGURANCA DO HOMELAB")
        set_font(left, size=8.3, color=MUTED, bold=True)
        p.add_run("\t")
        right = p.add_run("LEMBRADO")
        set_font(right, size=8.3, color=MUTED, bold=True)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Uso interno - sem segredos | Pagina ")
    set_font(r, size=8.3, color=MUTED)
    add_field(p, "PAGE")


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("AUDITORIA TECNICA")
    set_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(44)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Seguranca, resiliencia e operacao do homelab")
    set_font(r, size=27, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Ambiente de producao do sistema Lembrado")
    set_font(r, size=14, color=MUTED)

    add_callout(
        doc,
        "Conclusao executiva:",
        "o ambiente esta operacionalmente estavel, mas apresenta lacunas relevantes de recuperacao e seguranca. A prioridade absoluta e estabelecer backup verificavel antes de qualquer atualizacao, fechamento de portas ou alteracao de infraestrutura.",
        "warning",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("ESCOPO")
    set_font(r, size=9, color=MUTED, bold=True)
    add_paragraph(doc, "Servidor 192.168.1.11, containers Docker, rede, proxy, tunel, persistencia, atualizacoes, autenticacao administrativa e exposicao publica.", after=12)

    metadata = [
        ("Data da avaliacao", "29 de julho de 2026"),
        ("Metodo", "Levantamento remoto somente leitura"),
        ("Classificacao", "Uso interno - contem informacoes de infraestrutura, sem credenciais"),
        ("Responsavel pelo ambiente", "Leandro / Lembrado"),
    ]
    add_table(doc, ["Campo", "Informacao"], metadata, [2250, 7110], font_size=9.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Nenhum segredo, token, chave de API, senha, chave SSH ou valor de variavel sensivel foi incluido neste documento.")
    set_font(r, size=9.3, color=RED, bold=True)


def build_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    add_custom_numbering(doc)
    setup_header_footer(doc.sections[0], first=True)

    core = doc.core_properties
    core.title = "Auditoria de seguranca, resiliencia e operacao do homelab"
    core.subject = "Ambiente de producao do sistema Lembrado"
    core.creator = ""
    core.last_modified_by = ""
    core.keywords = "homelab; seguranca; Docker; backup; resiliencia"
    core.comments = "Uso interno. Documento sem credenciais ou valores secretos."

    add_cover(doc)
    doc.add_page_break()

    add_heading(doc, "1. Resumo executivo", 1)
    add_paragraph(
        doc,
        "A auditoria identificou um ambiente enxuto, funcional e com sinais positivos de estabilidade: todos os nove containers estavam ativos, sem reinicializacoes, o sistema respondia rapidamente, os endpoints sensiveis testados exigiam autenticacao e os cabecalhos publicos de seguranca estavam presentes.",
    )
    add_paragraph(
        doc,
        "O risco principal nao esta em uma indisponibilidade atual, mas na combinacao de ausencia de backup restauravel, superficies administrativas expostas na rede, componentes com privilegios equivalentes a root e um proxy reverso atrasado em relacao a correcoes de seguranca. Essas lacunas podem transformar uma falha operacional, credencial comprometida ou atualizacao mal-sucedida em perda prolongada do servico.",
    )
    add_callout(
        doc,
        "Prioridade obrigatoria:",
        "nao atualizar containers, alterar firewall, excluir DNS, remover volumes ou fechar portas antes de criar backups fora do host e validar ao menos uma restauracao em ambiente isolado.",
        "risk",
    )

    add_heading(doc, "1.1 Classificacao geral", 2)
    add_table(
        doc,
        ["Dimensao", "Avaliacao", "Leitura executiva"],
        [
            ("Disponibilidade atual", "Boa", "Containers estaveis e respostas HTTP rapidas."),
            ("Recuperacao de desastre", "Critica", "Nao foi encontrada rotina de backup ou teste de restauracao."),
            ("Superficie de ataque", "Alta", "Portas administrativas e da aplicacao escutam em todas as interfaces."),
            ("Gestao de privilegios", "Alta", "Docker, Portainer e socket Docker ampliam o impacto de uma credencial comprometida."),
            ("Atualizacoes", "Moderada/alta", "Nginx Proxy Manager requer prioridade; demais componentes exigem janela planejada."),
            ("Observabilidade", "Moderada", "Logs existem, mas faltam healthchecks, limites e alertas estruturados."),
        ],
        [2160, 1440, 5760],
        font_size=9.0,
    )

    add_heading(doc, "1.2 Ordem recomendada", 2)
    for text in [
        "Backup completo e teste de restauracao.",
        "Confirmacao da titularidade do IP antigo e limpeza de DNS, proxies e certificados obsoletos.",
        "Atualizacao controlada do Nginx Proxy Manager.",
        "Reducao da exposicao de portas e implantacao de politica de firewall.",
        "Hardening do SSH e separacao de privilegios administrativos.",
        "Protecao e rotacao planejada de segredos.",
        "Healthchecks, limites de recursos, rotacao de logs e monitoramento.",
        "Atualizacoes restantes e plano de migracao do Debian.",
        "Rotina continua de testes de restauracao e revisoes trimestrais.",
    ]:
        add_numbered(doc, text)

    add_heading(doc, "2. Escopo, metodo e limitacoes", 1)
    add_paragraph(
        doc,
        "A avaliacao foi executada por acesso SSH autorizado ao host, com comandos somente leitura. Foram inspecionados sistema operacional, capacidade, processos, portas, containers, imagens, redes, volumes, configuracoes expostas pelo Docker, logs recentes, DNS publico, certificados e respostas HTTP.",
    )
    add_heading(doc, "2.1 Itens verificados", 2)
    for item in [
        "Inventario de hardware, sistema operacional, kernel, memoria, disco, swap, uptime e sincronizacao de horario.",
        "Servicos escutando no host e portas publicadas por containers.",
        "Versoes, usuarios efetivos, capacidades, perfis de seguranca, reinicializacoes e consumo dos containers.",
        "Persistencia local, volumes Docker, configuracoes de stacks e existencia de rotinas de backup.",
        "Logs das ultimas 24 horas, renovacao de certificados e estado do erro 502 anteriormente investigado.",
        "Resolucao DNS, caminho do Cloudflare Tunnel, cabecalhos HTTP e protecao basica de endpoints internos.",
        "Estado de atualizacoes do Debian e comparacao de componentes criticos com versoes publicadas.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.2 Limitacoes", 2)
    for item in [
        "As regras efetivas de nftables e encaminhamentos do roteador nao foram confirmados porque a elevacao sudo exige senha interativa.",
        "Nao foi realizado teste de restauracao, varredura ativa, exploracao, alteracao de configuracao ou reinicializacao.",
        "A titularidade do endereco 2.57.91.92, ainda referenciado por dominios antigos, permanece pendente de confirmacao do responsavel.",
        "Valores de segredos e credenciais nao foram coletados nem reproduzidos; somente os nomes das categorias de configuracao foram considerados.",
        "Avaliacao pontual: o estado pode mudar apos atualizacoes, alteracoes de DNS, novos containers ou mudancas no roteador.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Arquitetura observada", 1)
    add_callout(
        doc,
        "Fluxo principal:",
        "Internet -> Cloudflare Tunnel -> cloudflared -> 192.168.1.11:80 -> Nginx Proxy Manager -> gestor_app:3000.",
        "info",
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.25
    diagram = (
        "Internet\n"
        "  -> Cloudflare Tunnel\n"
        "     -> cloudflared\n"
        "        -> host 192.168.1.11:80\n"
        "           -> Nginx Proxy Manager\n"
        "              -> gestor_app:3000\n\n"
        "Rede Docker gestor_gestor-network\n"
        "  gestor_app | gestor_worker | gestor_scheduler\n"
        "  evolution_api | evolution_db (PostgreSQL) | redis_queue | nginx-npm-1\n\n"
        "Rede bridge\n"
        "  cloudflared | portainer\n"
        "  Portainer -> /var/run/docker.sock (leitura e escrita)"
    )
    r = p.add_run(diagram)
    set_font(r, name="Consolas", size=9.4, color=DARK_BLUE)

    add_heading(doc, "3.1 Dependencias de dados", 2)
    add_bullet(doc, "Banco principal da aplicacao: Supabase externo.")
    add_bullet(doc, "Banco PostgreSQL local: dados da Evolution API.")
    add_bullet(doc, "Redis local: filas e estado operacional.")
    add_bullet(doc, "Volumes locais: sessoes da Evolution, dados do Nginx Proxy Manager, certificados e Portainer.")
    add_callout(
        doc,
        "Ponto de falha:",
        "todos os dados locais persistentes identificados residem no mesmo disco fisico do host. A perda do equipamento ou do sistema de arquivos pode atingir simultaneamente dados e configuracoes.",
        "risk",
    )

    add_heading(doc, "4. Inventario tecnico", 1)
    add_heading(doc, "4.1 Host", 2)
    add_table(
        doc,
        ["Item", "Estado observado"],
        [
            ("Hardware", "Lenovo ideapad 320-14IKB; arquitetura x86-64; 4 CPUs logicas."),
            ("Sistema", "Debian GNU/Linux 12 (bookworm), kernel 6.1.0-48-amd64."),
            ("Memoria", "3,6 GiB totais; aproximadamente 2,0 GiB em uso; 1,6 GiB disponiveis."),
            ("Swap", "975 MiB totais; aproximadamente 240 MiB em uso."),
            ("Armazenamento", "167,7 GB; raiz ext4 com cerca de 17 GB usados e 138 GB livres (11%)."),
            ("Operacao", "Uptime aproximado de 62 dias e 23 horas; carga baixa; NTP sincronizado."),
            ("Energia", "Bateria informada como cheia, 100%."),
            ("Fuso horario", "America/Sao_Paulo."),
        ],
        [2250, 7110],
        font_size=9.3,
    )

    add_heading(doc, "4.2 Containers e versoes", 2)
    add_table(
        doc,
        ["Componente", "Imagem/versao", "Estado"],
        [
            ("Aplicacao", "ghcr.io/leoaraujof/gestormaster:latest; commit f4d5fbdb94db", "Ativo; UID 1001"),
            ("Worker", "mesma imagem da aplicacao", "Ativo; UID 1001"),
            ("Scheduler", "mesma imagem da aplicacao", "Ativo; UID 1001"),
            ("Evolution API", "evoapicloud/evolution-api:v2.3.7", "Ativo; root"),
            ("PostgreSQL", "postgres:15-alpine; 15.18", "Ativo; UID 70"),
            ("Redis", "redis:7-alpine; 7.4.9", "Ativo; UID 999"),
            ("Cloudflared", "latest; 2026.7.2", "Ativo; patch 2026.7.3 sugerido"),
            ("Nginx Proxy Manager", "latest; 2.14.0", "Ativo; root; atualizar"),
            ("Portainer CE", "latest; 2.39.2", "Ativo; root"),
        ],
        [2520, 4320, 2520],
        font_size=8.8,
    )
    add_paragraph(doc, "Todos os nove containers estavam ativos, com politica de reinicializacao configurada e zero reinicializacoes observadas.", italic=True, color=MUTED, after=10)

    add_heading(doc, "4.3 Portas escutando no host", 2)
    add_table(
        doc,
        ["Porta", "Servico", "Exposicao observada"],
        [
            ("22", "SSH", "Todas as interfaces IPv4/IPv6"),
            ("80/443", "Nginx Proxy Manager", "Todas as interfaces IPv4/IPv6"),
            ("81", "Administracao NPM", "Todas as interfaces IPv4/IPv6"),
            ("3000", "gestor_app", "Todas as interfaces IPv4/IPv6"),
            ("8000/9443", "Portainer", "Todas as interfaces IPv4/IPv6"),
            ("8080", "Evolution API", "Somente rede interna Docker"),
            ("5432", "PostgreSQL", "Somente rede interna Docker"),
            ("6379", "Redis", "Somente rede interna Docker"),
        ],
        [1200, 3000, 5160],
        font_size=9.2,
    )
    add_paragraph(doc, "Nao foi encontrado UFW instalado. O pacote nftables existe, mas as regras efetivas nao foram auditadas sem elevacao sudo.", italic=True, color=MUTED)

    add_heading(doc, "4.4 Volumes persistentes", 2)
    for item in [
        "gestor_evolution_db_data - banco PostgreSQL da Evolution.",
        "gestor_evolution_instances - sessoes e instancias da Evolution.",
        "gestor_redis_queue_data - dados persistentes do Redis.",
        "nginx_npm_data e nginx_npm_letsencrypt - configuracoes, estado e certificados do proxy.",
        "portainer_data - configuracao do Portainer e arquivos das stacks.",
        "portainer_date - volume possivelmente incorreto ou sem uso; deve ser verificado antes de qualquer remocao.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. Controles positivos encontrados", 1)
    for item in [
        "Todos os containers estavam em execucao e sem reinicializacoes observadas.",
        "Aplicacao, worker e scheduler executam com usuario nao root (UID 1001).",
        "Nenhum container foi identificado como privileged ou com cap-add explicito.",
        "Perfis padrao de AppArmor e seccomp estavam ativos.",
        "Evolution API e painel de filas rejeitaram acesso sem autenticacao.",
        "O site publico respondeu com HSTS, CSP, protecao contra framing, nosniff, politica de permissoes e politica de referencia.",
        "Cloudflare Tunnel reduz a necessidade de publicar diretamente o IP residencial para o dominio principal.",
        "Pacotes Debian nao apresentavam atualizacoes pendentes no momento da coleta; timers diarios estavam funcionando.",
        "Disco com ampla folga e carga do host baixa.",
        "Nao foram encontrados novos registros de 'upstream sent too big header' apos o ajuste de buffers, indicando sucesso da correcao do 502.",
    ]:
        add_bullet(doc, item)
    add_callout(
        doc,
        "Leitura correta:",
        "os controles existentes reduzem riscos, mas nao substituem backup restauravel, segmentacao administrativa, politica de firewall e gestao de privilegios.",
        "positive",
    )

    add_heading(doc, "6. Registro de achados", 1)
    findings = [
        ("F-01", "Alta", "Ausencia de backup e restauracao", "Imediata"),
        ("F-02", "Alta", "Nginx Proxy Manager 2.14.0 atrasado em correcoes de seguranca", "Apos backup"),
        ("F-03", "Alta", "Portas administrativas e aplicacao expostas em todas as interfaces", "Curto prazo"),
        ("F-04", "Alta", "Portainer e grupo docker com privilegio equivalente a root", "Curto prazo"),
        ("F-05", "Alta", "Segredos de alto valor disponiveis a administradores Docker/Portainer", "Curto prazo"),
        ("F-06", "Alta operacional", "DNS antigo e renovacoes ACME falhando repetidamente", "Apos confirmacao"),
        ("F-07", "Media/alta", "SSH permite superficie maior que a necessaria", "Curto prazo"),
        ("F-08", "Media", "Sem healthchecks e limites de recursos", "Medio prazo"),
        ("F-09", "Media", "Logs Docker sem rotacao explicita", "Medio prazo"),
        ("F-10", "Media", "Redis com vm.overcommit_memory=0", "Curto prazo"),
        ("F-11", "Media", "Politica de atualizacoes e ciclo de vida do Debian", "Planejado"),
        ("F-12", "Baixa", "Exposicao de tecnologia e CSP permissiva", "Oportunistica"),
    ]
    add_table(doc, ["ID", "Severidade", "Achado", "Tratamento"], findings, [900, 1450, 5160, 1850], font_size=8.7)

    detailed = [
        (
            "6.1 F-01 - Ausencia de backup e restauracao",
            "Alta",
            "Nao foram encontrados restic, borg, rclone, duplicity, rsnapshot, rsync, diretorios de backup, dumps SQL, crontab do usuario ou servico dedicado. Os unicos backups identificados sao rotinas de banco de pacotes do Debian, que nao protegem dados da aplicacao.",
            "Falha de disco, erro humano, corrupcao, atualizacao malsucedida ou comprometimento podem causar perda simultanea de banco local, sessoes, certificados, configuracoes do proxy e stacks.",
            "Implementar backup criptografado fora do host, retencao definida, monitoramento de sucesso e teste de restauracao isolado. Exportar tambem os arquivos Compose armazenados dentro do volume do Portainer.",
        ),
        (
            "6.2 F-02 - Nginx Proxy Manager atrasado",
            "Alta",
            "O ambiente executa a versao 2.14.0. A versao 2.15.0 inclui OpenResty atualizado e correcoes associadas a vulnerabilidades publicadas.",
            "O proxy esta na fronteira de entrada HTTP/HTTPS. Uma vulnerabilidade nessa camada pode afetar confidencialidade, integridade e disponibilidade.",
            "Atualizar somente depois do backup e de um plano de rollback. Fixar uma versao explicita, validar banco/configuracao e executar testes de rotas e certificados.",
        ),
        (
            "6.3 F-03 - Exposicao excessiva de portas",
            "Alta",
            "As portas 81, 3000, 8000 e 9443 escutam em todas as interfaces, alem das portas esperadas 22, 80 e 443.",
            "Qualquer dispositivo na rede local, e possivelmente regras de encaminhamento do roteador, pode alcancar interfaces administrativas ou contornar o proxy.",
            "Remover a publicacao direta da porta 3000, restringir administracao a localhost, VPN, Cloudflare Access ou sub-rede administrativa, remover 8000 se o Edge Agent nao for usado e adotar firewall default-deny.",
        ),
        (
            "6.4 F-04 - Privilegios equivalentes a root",
            "Alta",
            "Portainer monta /var/run/docker.sock com leitura e escrita; o usuario leandro pertence ao grupo docker. Portainer, NPM e Evolution executam como root no container.",
            "Acesso ao socket ou ao grupo docker permite controle amplo do host. Uma conta administrativa comprometida pode criar containers privilegiados, montar o sistema de arquivos e extrair segredos.",
            "Minimizar administradores, proteger Portainer com MFA/controle de acesso, avaliar socket-proxy com permissoes minimas, separar usuario de deploy e revisar a necessidade de root em cada container.",
        ),
        (
            "6.5 F-05 - Segredos em variaveis de ambiente",
            "Alta",
            "Foram identificadas categorias sensiveis como service role do Supabase, segredos Stripe, chave Evolution, chave OpenAI, senha Redis, chave de criptografia, pepper de PIN, segredos PixGo e segredo de cron. Nenhum valor foi registrado.",
            "Administradores Docker/Portainer e processos com acesso equivalente ao daemon podem inspecionar esses valores.",
            "Migrar segredos de maior impacto para Docker Compose secrets ou cofre criptografado, limitar Portainer e definir procedimento de rotacao. Rotacionar imediatamente se houver suspeita de exposicao.",
        ),
        (
            "6.6 F-06 - DNS antigo e certificados falhando",
            "Alta operacional",
            "Dominios roboajuda.site ainda resolvem para 2.57.91.92, enquanto rotas antigas permanecem no NPM. As renovacoes Let's Encrypt falham repetidamente porque o desafio ACME chega ao endereco antigo.",
            "Alertas recorrentes, certificados proximos do vencimento e risco de trafego ou credenciais serem direcionados a infraestrutura fora do controle atual.",
            "Confirmar a titularidade de 2.57.91.92. Se nao pertencer mais ao responsavel, remover ou corrigir DNS imediatamente, desativar proxies obsoletos e revogar/remover certificados sem uso.",
        ),
        (
            "6.7 F-07 - Hardening do SSH",
            "Media/alta",
            "O handshake anuncia publickey e password; X11Forwarding esta habilitado e fail2ban nao foi encontrado. KbdInteractiveAuthentication esta desabilitado.",
            "Aumento da superficie para forca bruta, uso indevido de senha e recursos desnecessarios.",
            "Depois de validar duas sessoes por chave, desabilitar senha e root, desabilitar X11, limitar usuarios e tentativas. Restringir a chave de automacao por origem e opcoes de authorized_keys.",
        ),
        (
            "6.8 F-08 - Ausencia de healthchecks e limites",
            "Media",
            "Nenhum container possui healthcheck. Nao ha limites de memoria, CPU ou PIDs. O worker consumia aproximadamente 667 MiB, cerca de 18% da RAM.",
            "Um processo travado pode continuar marcado como ativo; crescimento descontrolado pode pressionar memoria e afetar todos os servicos.",
            "Adicionar healthchecks adequados, heartbeat de worker/scheduler e limites graduais baseados em medicoes, com alertas antes de reinicios automaticos.",
        ),
        (
            "6.9 F-09 - Logs sem rotacao",
            "Media",
            "O driver json-file esta em uso sem max-size e max-file configurados.",
            "Crescimento indefinido dos logs pode ocupar o disco e causar indisponibilidade.",
            "Configurar rotacao, por exemplo max-size 10m e max-file entre 3 e 5, validando necessidades de retencao e suporte.",
        ),
        (
            "6.10 F-10 - Ajuste de memoria do Redis",
            "Media",
            "O Redis alerta que vm.overcommit_memory esta desabilitado; o host reporta valor 0.",
            "Operacoes de fork, persistencia ou baixa memoria podem falhar em cenarios de pressao.",
            "Definir vm.overcommit_memory=1 de forma persistente em janela controlada e validar o alerta apos a alteracao.",
        ),
        (
            "6.11 F-11 - Atualizacoes e ciclo de vida",
            "Media",
            "Debian 12 entrou em LTS; unattended-upgrades nao esta instalado. Docker 29.5.2 estava atras do ramo 29.6.2 e cloudflared 2026.7.2 informava patch 2026.7.3.",
            "Atraso acumulado aumenta exposicao e torna futuras atualizacoes mais arriscadas.",
            "Definir calendario de manutencao, atualizar por criticidade, fixar imagens por versao ou digest e planejar migracao para Debian 13 em 30 a 90 dias, com teste de drivers e rollback.",
        ),
        (
            "6.12 F-12 - Endurecimento HTTP adicional",
            "Baixa",
            "A CSP permite unsafe-inline para scripts e estilos, e o cabecalho X-Powered-By revela Next.js.",
            "Esses itens ajudam reconhecimento e reduzem a resistencia da CSP, embora os demais cabecalhos estejam bem configurados.",
            "Remover X-Powered-By e evoluir a CSP para nonce/hash quando compativel com a aplicacao, sem quebrar funcionalidades.",
        ),
    ]
    for title, severity, evidence, impact, recommendation in detailed:
        add_heading(doc, title, 2)
        kind = "risk" if severity.startswith("Alta") else "warning" if severity.startswith("Media") else "info"
        add_callout(doc, "Severidade:", severity, kind)
        add_paragraph(doc, "Evidencia. " + evidence, bold_prefix="Evidencia. ")
        add_paragraph(doc, "Impacto. " + impact, bold_prefix="Impacto. ")
        add_paragraph(doc, "Recomendacao. " + recommendation, bold_prefix="Recomendacao. ")

    add_heading(doc, "7. Evidencias operacionais relevantes", 1)
    add_heading(doc, "7.1 Estado do erro 502", 2)
    add_paragraph(
        doc,
        "O erro anterior 'upstream sent too big header while reading response header from upstream' teve ultimo registro em 29/07/2026 as 12:04. Nao foram encontrados novos eventos depois do aumento dos buffers do proxy.",
    )
    add_callout(doc, "Resultado:", "a correcao de buffers foi bem-sucedida nas evidencias observadas. O problema deve continuar monitorado por pelo menos sete dias.", "positive")

    add_heading(doc, "7.2 Logs e disponibilidade", 2)
    add_bullet(doc, "Aplicacao respondeu HTTP 200 em aproximadamente 9 ms no teste local.")
    add_bullet(doc, "Nginx Proxy Manager respondeu HTTP 200 em aproximadamente 15 ms.")
    add_bullet(doc, "Nao foram encontrados padroes de erro nos logs principais da aplicacao nas ultimas 24 horas.")
    add_bullet(doc, "Houve um erro transitorio de stream WhatsApp com codigo 503 na Evolution.")
    add_bullet(doc, "Duas mensagens 'role root does not exist' no PostgreSQL foram produzidas pelo comando de auditoria pg_isready sem usuario explicito, nao por falha da aplicacao.")
    add_bullet(doc, "As falhas mais numerosas do NPM eram renovacoes ACME repetidas dos dominios antigos.")

    doc.add_page_break()
    add_heading(doc, "7.3 DNS e tunel", 2)
    add_table(
        doc,
        ["Nome", "Destino/estado"],
        [
            ("lembrado.com.br e www", "Cloudflare; site funcional pelo Tunnel."),
            ("Origem do Tunnel", "http://192.168.1.11:80."),
            ("api.roboajuda.site", "2.57.91.92; configuracao antiga."),
            ("queue.roboajuda.site", "2.57.91.92; configuracao antiga."),
            ("portainer.roboajuda.site", "2.57.91.92; configuracao antiga."),
            ("proxy.roboajuda.site", "2.57.91.92; configuracao antiga."),
            ("roboajuda.site", "2.57.91.92; configuracao antiga."),
        ],
        [2880, 6480],
        font_size=9.2,
    )
    add_paragraph(
        doc,
        "A aplicacao usa EVOLUTION_API_URL interno, apontando para o servico Docker evolution-api:8080. Portanto, nas configuracoes observadas, a chave Evolution nao e enviada ao IP antigo.",
        italic=True,
        color=MUTED,
    )

    add_heading(doc, "8. Plano de remediacao recomendado", 1)
    add_callout(
        doc,
        "Regra de mudanca:",
        "cada fase deve ter backup, criterio de aceite, rollback documentado e janela de manutencao. Alteracoes de rede devem preservar uma sessao SSH de contingencia ate a validacao final.",
        "warning",
    )
    phases = [
        (
            "Fase 0 - Recuperacao antes de qualquer mudanca",
            [
                "Exportar as stacks e arquivos Compose do Portainer.",
                "Executar pg_dump consistente do PostgreSQL da Evolution.",
                "Copiar o volume de instancias/sessoes da Evolution com o servico em estado consistente.",
                "Salvar nginx_npm_data, nginx_npm_letsencrypt e portainer_data.",
                "Manter copia criptografada fora do host, com retencao definida.",
                "Restaurar em ambiente temporario isolado e registrar evidencias do teste.",
            ],
            "Backup concluido, hash/manifesto registrado e restauracao funcional comprovada.",
        ),
        (
            "Fase 1 - DNS e certificados antigos",
            [
                "Confirmar quem controla 2.57.91.92.",
                "Inventariar dependencias ainda ativas de roboajuda.site.",
                "Corrigir ou remover DNS antigo, proxies e certificados sem uso.",
                "Confirmar que as renovacoes ACME deixam de falhar.",
            ],
            "Nenhum dominio controlado aponta para infraestrutura indevida; logs ACME limpos.",
        ),
        (
            "Fase 2 - Atualizacao do proxy",
            [
                "Fixar Nginx Proxy Manager em versao 2.15.0 ou superior validada.",
                "Registrar imagem anterior e procedimento de rollback.",
                "Atualizar em janela; testar login, rotas, certificados e cabecalhos.",
            ],
            "Site, APIs e painel respondem; certificados validos; sem regressao de 502.",
        ),
        (
            "Fase 3 - Reducao da superficie de rede",
            [
                "Remover publicacao direta da aplicacao na porta 3000 quando o NPM compartilhar a rede Docker.",
                "Restringir 81 e 9443 a origem administrativa segura.",
                "Remover 8000 se Edge Agent nao estiver em uso.",
                "Revisar encaminhamentos do roteador e implantar firewall default-deny.",
                "Considerar cloudflared conectado diretamente a rede Docker e ao nome do servico.",
            ],
            "Somente portas justificadas permanecem acessiveis; administracao funciona apenas pelo caminho autorizado.",
        ),
        (
            "Fase 4 - SSH e privilegios",
            [
                "Validar login por chave em duas sessoes antes de desabilitar senha.",
                "Definir PasswordAuthentication no, PermitRootLogin no e X11Forwarding no.",
                "Aplicar AllowUsers e MaxAuthTries restritivos.",
                "Restringir chave de automacao por origem e opcoes restrict.",
                "Revisar membros do grupo docker e administradores do Portainer.",
            ],
            "Login administrativo por chave funciona; senha/root/X11 bloqueados; sem perda de acesso.",
        ),
        (
            "Fase 5 - Segredos",
            [
                "Classificar segredos por impacto e dependencia.",
                "Migrar os mais criticos para secrets/cofre criptografado.",
                "Limitar quem pode inspecionar containers.",
                "Rotacionar segredos com procedimento coordenado e teste de rollback.",
            ],
            "Valores nao aparecem em configuracoes acessiveis alem do necessario; aplicacao e webhooks continuam validos.",
        ),
        (
            "Fase 6 - Resiliencia e observabilidade",
            [
                "Definir vm.overcommit_memory=1.",
                "Adicionar healthchecks e heartbeat para jobs.",
                "Aplicar limites de memoria, CPU e PIDs gradualmente.",
                "Configurar rotacao de logs Docker.",
                "Alertar falhas de backup, expiracao de certificados, disco, memoria e containers.",
            ],
            "Alertas testados, healthchecks confiaveis, logs limitados e recursos sem regressao.",
        ),
        (
            "Fase 7 - Atualizacoes e ciclo continuo",
            [
                "Atualizar cloudflared e Docker em janelas separadas.",
                "Adotar unattended-upgrades ou calendario formal de patches.",
                "Planejar Debian 13 com teste de compatibilidade e rollback.",
                "Substituir latest por versoes ou digests controlados.",
                "Executar teste de restauracao recorrente e revisao trimestral.",
            ],
            "Inventario versionado, rotina documentada e evidencias periodicas de recuperacao.",
        ),
    ]
    for idx, (title, actions, acceptance) in enumerate(phases):
        add_heading(doc, title, 2)
        for action in actions:
            add_bullet(doc, action)
        add_callout(doc, "Criterio de aceite:", acceptance, "positive")

    add_heading(doc, "9. Checklist de backup e restauracao", 1)
    checklist = [
        ("Inventario", "Volumes, stacks, imagens, versoes e dependencias documentados."),
        ("PostgreSQL Evolution", "Dump logico validado e restaurado em instancia isolada."),
        ("Sessoes Evolution", "Volume copiado em estado consistente; permissoes preservadas."),
        ("NPM", "Dados, certificados e banco/configuracao recuperaveis."),
        ("Portainer", "portainer_data e arquivos Compose exportados fora do volume."),
        ("Redis", "Decidir se estado precisa de backup; validar AOF/RDB conforme uso."),
        ("Segredos", "Backup criptografado; acesso e chave de recuperacao separados."),
        ("Copia externa", "Destino fora do host; criptografia; retencao e imutabilidade quando possivel."),
        ("Teste", "Restauracao executada, tempo medido e evidencias salvas."),
        ("Monitoramento", "Falha, atraso ou tamanho anormal do backup gera alerta."),
    ]
    add_table(doc, ["Controle", "Condicao minima"], checklist, [2520, 6840], font_size=9.2)

    add_heading(doc, "10. Criterios de validacao final", 1)
    for item in [
        "Backup externo existe, pode ser lido e foi restaurado em ambiente isolado.",
        "Site lembrado.com.br, login, painel, renovacoes, PIX, WhatsApp e jobs criticos passam por smoke test.",
        "Nenhum novo 502 ou 'too big header' aparece durante a janela de observacao.",
        "Administracao NPM/Portainer nao e acessivel de redes nao autorizadas.",
        "Porta 3000 nao contorna o proxy; portas 8000/9443/81 seguem a politica definida.",
        "SSH por chave funciona e autenticacao por senha/root esta bloqueada.",
        "NPM esta atualizado e com rollback conhecido.",
        "Renovacoes ACME funcionam ou configuracoes antigas foram removidas.",
        "Healthchecks refletem o estado real; alertas e rotacao de logs foram testados.",
        "Segredos permanecem funcionais, nao foram registrados em logs e possuem responsavel de rotacao.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "11. Fontes tecnicas", 1)
    sources = [
        ("Docker - privilegios do grupo docker", "https://docs.docker.com/engine/install/linux-postinstall/"),
        ("Docker - seguranca do daemon", "https://docs.docker.com/engine/security/"),
        ("Docker - rotacao do driver json-file", "https://docs.docker.com/engine/logging/drivers/json-file/"),
        ("Cloudflare - conectividade por Tunnel", "https://developers.cloudflare.com/cloudflare-one/networks/connectivity-options/"),
        ("Nginx Proxy Manager - releases", "https://github.com/NginxProxyManager/nginx-proxy-manager/releases"),
        ("Portainer - releases", "https://github.com/portainer/portainer/releases"),
        ("Evolution API - releases", "https://github.com/evolution-foundation/evolution-api/releases"),
        ("Docker Engine 29 - notas de versao", "https://docs.docker.com/engine/release-notes/29/"),
        ("Debian 12 LTS", "https://www.debian.org/News/2026/20260712"),
        ("Debian - sshd_config", "https://manpages.debian.org/bookworm/openssh-server/sshd_config.5.en.html"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r1 = p.add_run(label + ": ")
        set_font(r1, size=9.5, color=INK, bold=True)
        r2 = p.add_run(url)
        set_font(r2, size=9.5, color=BLUE)

    add_heading(doc, "12. Observacao de confidencialidade", 1)
    add_paragraph(
        doc,
        "Este relatorio contem enderecos de rede, nomes de servicos e arquitetura suficientes para orientar manutencao. Deve ser compartilhado somente com pessoas autorizadas. O documento foi deliberadamente produzido sem senhas, tokens, chaves, valores de variaveis de ambiente, fingerprints, IDs de maquina ou dados de historico de acesso.",
    )
    add_callout(
        doc,
        "Estado da auditoria:",
        "levantamento concluido sem alteracoes no homelab. A implementacao deve seguir a ordem do plano e exigir aprovacao explicita para cada mudanca com impacto.",
        "info",
    )

    for section in doc.sections:
        setup_page(section)

    apply_portuguese_diacritics(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
